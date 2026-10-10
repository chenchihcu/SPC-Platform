from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from app.data import master_data_db
from app.data.paste_printing_spec_registry import save as save_paste_spec
from app.data.session_store import SessionStore
from app.data.stencil_document_library import create_stencil_document_version
from app.data.stencil_thickness_registry import save as save_stencil_spec
from app.services.analysis_orchestrator import STATUS_READY, AnalysisOrchestrator
from app.services.diagnostic_excel_exporter import export_diagnostic_summary_xlsx
from app.services.report_context import build_pptx_report_context
from app.services.stencil_document_service import PARSER_VERSION
from app.utils.constants import FILTER_ALL, RANGE_ALL_BOARDS
from app.viewmodels.chart_analysis_viewmodel import compute_analysis_payload


@pytest.fixture()
def isolated_master_db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    db_file = tmp_path / "spc_master.db"
    monkeypatch.setenv("SPC_MASTER_DB_PATH", str(db_file))
    master_data_db._INITIALIZED = False
    master_data_db.init_db()
    yield db_file
    SessionStore().clear()
    master_data_db._INITIALIZED = False


def _create_active_nominal(tmp_path: Path, product: str = "P1") -> None:
    gerber = tmp_path / "paste.gtp"
    gerber.write_text(
        "G04 test*\n%FSLAX46Y46*%\n%MOMM*%\n%ADD10R,0.600000X0.300000*%\n"
        "D10*\nX0000000000Y0000000000D03*\nM02*\n",
        encoding="ascii",
    )
    pads = tmp_path / "pads.csv"
    pads.write_text(
        "RefDes,Pad,X,Y,Width,Height,Rotation,Side\nU1,01,0,0,0.6,0.3,0,TOP\n",
        encoding="utf-8",
    )
    assert save_paste_spec(
        {
            "product_name": product,
            "default_volume_target": 100.0,
            "default_volume_lsl": 70.0,
            "default_volume_usl": 150.0,
            "default_area_target": 100.0,
            "default_area_lsl": 70.0,
            "default_area_usl": 150.0,
            "default_height_lsl": 70.0,
            "default_height_usl": 140.0,
        }
    )
    assert save_stencil_spec(
        {
            "product_name": product,
            "stencil_type": "normal",
            "thickness_main": 0.06,
            "unit_mode": "absolute",
        }
    )
    create_stencil_document_version(
        product_name=product,
        stencil_no="ST-001",
        revision="A",
        side="TOP",
        gerber_source=str(gerber),
        pad_csv_source=str(pads),
        coordinate_transform={"rotation_deg": 0, "mirrored": False, "scale": 1.0},
        validation_summary={
            "total_pads": 1,
            "matched_pads": 1,
            "coverage": 1.0,
            "duplicate_pad_rows": 0,
            "ambiguous_openings": 0,
            "unclassified_extra_openings": 0,
            "missing_thickness_pads": 0,
            "requires_user_overlay_confirmation": True,
            "overlay_confirmed": True,
        },
        parser_version=PARSER_VERSION,
        nominals=[
            {
                "refdes": "U1",
                "pad": "01",
                "side": "TOP",
                "opening_count": 1,
                "nominal_area_mm2": 0.18,
                "resolved_thickness_mm": 0.06,
                "nominal_volume_mm3": 0.0108,
                "match_status": "MATCHED",
                "preview_geometry": {},
            }
        ],
        activate=True,
    )


def _measurement_rows(*, include_unmatched: bool = False) -> pd.DataFrame:
    count = 12
    frame = pd.DataFrame(
        {
            "BoardNo": [f"B{i:02d}" for i in range(count)],
            "RefDes": ["U1"] * count,
            "BaseRefDes": ["U1"] * count,
            "Pad": ["01"] * count,
            "PartType": ["IC"] * count,
            "Volume": [0.0102 + i * 0.0001 for i in range(count)],
            "Area": [0.17 + i * 0.0005 for i in range(count)],
            "Height": [0.057 + i * 0.0002 for i in range(count)],
            "Side": ["TOP"] * count,
        }
    )
    if include_unmatched:
        frame.loc[count - 1, ["RefDes", "BaseRefDes", "Pad"]] = ["U2", "U2", "02"]
    return frame


def _prepare(store: SessionStore):
    return AnalysisOrchestrator().prepare_refresh(
        store=store,
        selected_features=["Volume"],
        range_mode=RANGE_ALL_BOARDS,
        board_specify="",
        refdes=FILTER_ALL,
        part_type=FILTER_ALL,
        optional_filters={},
        manual_workorder_spec={},
    )


def test_orchestrator_uses_percent_domain_only_at_full_coverage(
    isolated_master_db: Path, tmp_path: Path
) -> None:
    _create_active_nominal(tmp_path)
    store = SessionStore()
    store.clear()
    store.workorder_master = {"product_name": "P1"}
    store.meas_meta = {
        "is_valid": True,
        "side": "TOP",
        "measurement_units": {"Volume": "mm", "Area": "mm", "Height": "mm"},
    }
    store.meas_df = _measurement_rows()

    prep = _prepare(store)

    assert prep.status == STATUS_READY
    assert prep.analysis_metadata["value_domain"] == "percent"
    assert prep.analysis_metadata["coverage"] == 1.0
    assert prep.filtered_df is not None
    assert prep.filtered_df["Volume"].iloc[0] == pytest.approx(94.444444)
    assert prep.workorder_spec["height"]["target"] == "100.0"
    assert prep.usl == 150.0


def test_incomplete_coverage_keeps_raw_trends_and_blocks_spec_statistics(
    isolated_master_db: Path, tmp_path: Path
) -> None:
    _create_active_nominal(tmp_path)
    store = SessionStore()
    store.clear()
    store.workorder_master = {"product_name": "P1"}
    store.meas_meta = {
        "is_valid": True,
        "side": "TOP",
        "measurement_units": {"Volume": "mm", "Area": "mm", "Height": "mm"},
    }
    store.meas_df = _measurement_rows(include_unmatched=True)
    prep = _prepare(store)
    assert prep.status == STATUS_READY
    assert prep.analysis_metadata["value_domain"] == "raw"
    assert prep.analysis_metadata["verify"] is True
    assert prep.workorder_spec == {}
    assert prep.usl is None and prep.lsl is None

    payload, error = compute_analysis_payload(
        prep.filtered_df,
        prep.selected_features,
        prep.usl,
        prep.lsl,
        prep.target,
        prep.workorder_spec,
        workorder_master={"product_name": "P1", "_stencil_analysis": prep.analysis_metadata},
    )
    assert error is None and payload is not None
    assert payload["spc"]["metadata"]["is_valid"] is True
    assert payload["cap"]["metadata"]["is_valid"] is False
    assert payload["dist"]["metadata"]["is_valid"] is False
    assert payload["cap"]["metadata"]["verify"] is True
    assert "VERIFY" in payload["cap"]["metadata"]["error"]
    assert payload["summary"]["process"]["overall_yield_pct"] is None
    assert payload["stencil_analysis"]["matched"] == 11
    assert "U2/02" in payload["stencil_analysis"]["unmatched_pads"]


def test_report_context_and_xlsx_disclose_stencil_domain(
    isolated_master_db: Path, tmp_path: Path
) -> None:
    store = SessionStore()
    store.clear()
    store.workorder_master = {"product_name": "P1"}
    metadata = {
        "value_domain": "raw",
        "stencil_revision": "A",
        "stencil_no": "ST-001",
        "side": "TOP",
        "matched": 11,
        "total": 12,
        "coverage": 11 / 12,
        "verify": True,
        "verify_reasons": ["Pad 覆蓋不足：11/12"],
        "unmatched_pads": ["U2/02"],
        "original_units": {"Volume": "mm", "Area": "mm", "Height": "mm"},
    }
    store.analysis_domain_meta = metadata
    data = _measurement_rows(include_unmatched=True)
    summary = {"process": {"dashboard_layers": {}}, "per_measure": {}}
    payload = {"selected_features": ["Volume"], "summary": summary, "stencil_analysis": metadata}
    context = build_pptx_report_context(
        store=store,
        filtered_df=data,
        summary_data=summary,
        diagnostics=[],
        selected_features=["Volume"],
        batch=RANGE_ALL_BOARDS,
        refdes_filter=FILTER_ALL,
        part_type=FILTER_ALL,
        product="P1",
        time_start=None,
        time_end=None,
        line=None,
        analysis_payload=payload,
    )
    assert context["stencil_analysis"]["stencil_revision"] == "A"
    assert context["data_scope"]["section_trust"]["statistics"].startswith("VERIFY")
    assert "Volume%" in context["metric_definitions"]

    xlsx = tmp_path / "stencil.xlsx"
    export_diagnostic_summary_xlsx(payload, xlsx)
    workbook = load_workbook(xlsx, data_only=True)
    assert "鋼板開窗" in workbook.sheetnames
    values = [workbook["鋼板開窗"].cell(row=row, column=2).value for row in range(2, 14)]
    assert "VERIFY" in values
    assert "A" in values
    assert "11/12" in values
